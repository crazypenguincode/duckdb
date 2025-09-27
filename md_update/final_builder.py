#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
最终版 Word 公式构建器
使用迭代式 OMML 解析，实现高质量公式渲染
"""

import os
import re
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.shared import OxmlElement, qn
from docx.oxml.ns import nsdecls
from docx.oxml.parser import parse_xml
from xml.sax.saxutils import escape

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

BASE_DIR = Path(__file__).parent.absolute()
IMAGES_DIR = BASE_DIR / "images"

def create_base_template():
    doc = Document()
    for paragraph in doc.paragraphs:
        p = paragraph._element
        p.getparent().remove(p)
    doc.styles['Normal'].font.name = 'Times New Roman'
    doc.styles['Normal'].font.size = Pt(12)
    return doc

def force_font_style(run, font_name='Times New Roman', font_size=12, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    rPr = run._element.get_or_add_rPr()
    for rFonts in rPr.xpath('.//w:rFonts'):
        rPr.remove(rFonts)
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    rPr.insert(0, rFonts)
    for sz in rPr.xpath('.//w:sz'):
        rPr.remove(sz)
    for szCs in rPr.xpath('.//w:szCs'):
        rPr.remove(szCs)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(font_size * 2))
    rPr.append(sz)
    szCs = OxmlElement('w:szCs')
    szCs.set(qn('w:val'), str(font_size * 2))
    rPr.append(szCs)

def add_text_with_style(paragraph, text, bold=False, italic=False, font_size=12):
    run = paragraph.add_run(text)
    force_font_style(run, 'Times New Roman', font_size, bold, italic)
    return run

def _tex_to_omml(s: str) -> str:
    """迭代式 LaTeX -> OMML 转换器"""
    s = s.strip()
    s = re.sub(r'\\text\{([^}]+)\}', r'\1', s)
    s = s.replace(r'\left', '')
    s = s.replace(r'\right', '')

    # 迭代处理，由内而外
    for _ in range(5): # 限制迭代次数以防死循环
        # 1. 最内层分数
        while True:
            match = re.search(r'\\frac\{([^{}]*?)\}\{([^{}]*?)\}', s)
            if not match: break
            num = f'<m:r><m:t>{escape(match.group(1))}</m:t></m:r>'
            den = f'<m:r><m:t>{escape(match.group(2))}</m:t></m:r>'
            omml = f'<m:f><m:num>{num}</m:num><m:den>{den}</m:den></m:f>'
            s = s.replace(match.group(0), omml, 1)
        
        # 2. 最内层上下标
        while True:
            match = re.search(r'(\w+)_\{([^{}]*?)\}', s)
            if not match: break
            base = f'<m:r><m:t>{escape(match.group(1))}</m:t></m:r>'
            sub = f'<m:r><m:t>{escape(match.group(2))}</m:t></m:r>'
            omml = f'<m:sSub><m:e>{base}</m:e><m:sub>{sub}</m:sub></m:sSub>'
            s = s.replace(match.group(0), omml, 1)
        
        while True:
            match = re.search(r'(\w+)\^\{([^{}]*?)\}', s)
            if not match: break
            base = f'<m:r><m:t>{escape(match.group(1))}</m:t></m:r>'
            sup = f'<m:r><m:t>{escape(match.group(2))}</m:t></m:r>'
            omml = f'<m:sSup><m:e>{base}</m:e><m:sup>{sup}</m:sup></m:sSup>'
            s = s.replace(match.group(0), omml, 1)

        # 3. 带括号的上下标
        while True:
            match = re.search(r'\(([^()]*?)\)\^\{([^{}]*?)\}', s)
            if not match: break
            base = f'<m:d><m:e><m:r><m:t>{escape(match.group(1))}</m:t></m:r></m:e></m:d>'
            sup = f'<m:r><m:t>{escape(match.group(2))}</m:t></m:r>'
            omml = f'<m:sSup><m:e>{base}</m:e><m:sup>{sup}</m:sup></m:sSup>'
            s = s.replace(match.group(0), omml, 1)

    # 4. 处理特殊函数和符号
    s = s.replace(r'\ln', 'ln')
    s = s.replace(r'\approx', '≈')
    s = re.sub(r'\\log_2', '<m:sSub><m:e><m:r><m:t>log</m:t></m:r></m:e><m:sub><m:r><m:t>2</m:t></m:r></m:sub>', s)
    s = re.sub(r'\\frac\s*(\S+)\s*(\S+)', r'<m:f><m:num><m:r><m:t>\1</m:t></m:r></m:num><m:den><m:r><m:t>\2</m:t></m:r></m:den></m:f>', s)
    s = re.sub(r'(\w+)\^(\w+)', r'<m:sSup><m:e><m:r><m:t>\1</m:t></m:r></m:e><m:sup><m:r><m:t>\2</m:t></m:r></m:sup>', s)

    # 5. 将剩余的纯文本包裹
    parts = re.split(r'(<m:.*?>)', s)
    result = ""
    for part in parts:
        if not part: continue
        if part.startswith('<m:'):
            result += part
        elif part.strip():
            result += f'<m:r><m:t>{escape(part)}</m:t></m:r>'
    
    return result if result.strip() else f'<m:r><m:t>{escape(s)}</m:t></m:r>'

def create_native_math_equation(paragraph, latex_formula):
    try:
        inner_omml = _tex_to_omml(latex_formula)
        
        p_element = paragraph._element
        p_parent = p_element.getparent()
        
        # 创建新的 w:p 元素
        new_p = OxmlElement('w:p')
        
        # 复制段落属性
        if p_element.pPr is not None:
            new_p.append(p_element.pPr)
        
        # 创建并添加 m:oMathPara 和 m:oMath
        oMathPara = OxmlElement('m:oMathPara')
        oMath = OxmlElement('m:oMath')
        
        # 解析 inner_omml 并添加到 oMath
        # 添加命名空间声明以帮助解析器
        inner_xml = f'<root xmlns:m="{M_NS}" xmlns:w="{W_NS}">{inner_omml}</root>'
        frag = parse_xml(inner_xml)
        for child in frag:
            oMath.append(child)
            
        oMathPara.append(oMath)
        new_p.append(oMathPara)
        
        # 插入新段落并删除旧的
        p_parent.insert(p_parent.index(p_element) + 1, new_p)
        p_parent.remove(p_element)
        
        return True
    except Exception as e:
        print(f"❌ 原生公式创建失败: {latex_formula} -> {e}")
        return False

def process_chapter(content, doc, chapter_num):
    lines = content.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        if line.startswith('$$') and line.endswith('$$'):
            formula = line[2:-2].strip()
            p = doc.add_paragraph() # Add a placeholder paragraph
            create_native_math_equation(p, formula)
        elif line.startswith('$$'):
            formula_lines = [line[2:]]
            j = i + 1
            while j < len(lines):
                if lines[j].strip().endswith('$$'):
                    formula_lines.append(lines[j].strip()[:-2])
                    break
                else:
                    formula_lines.append(lines[j])
                j += 1
            formula = '\n'.join(formula_lines).strip()
            p = doc.add_paragraph()
            create_native_math_equation(p, formula)
            i = j
        elif '$' in line:
            p = doc.add_paragraph()
            parts = re.split(r'\$([^$]+)\$', line)
            for idx, part in enumerate(parts):
                if idx % 2 == 1:
                    # For inline math, we can't just replace the paragraph
                    # This is more complex, for now, use enhanced unicode
                    run = p.add_run(part)
                    run.font.name = 'Cambria Math'
                else:
                    add_text_with_style(p, part)
        elif line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            title_text = line.lstrip('#').strip()
            add_heading_with_style(doc, title_text, level)
        elif line.startswith('**图'):
            # ... (image processing logic from previous versions)
            p = doc.add_paragraph()
            add_text_with_style(p, line, bold=True)
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        elif line:
            p = doc.add_paragraph()
            process_text_line(p, line)
        
        i += 1

def add_heading_with_style(doc, text, level=1):
    heading = doc.add_heading('', level=level)
    heading.clear()
    font_sizes = {1: 18, 2: 16, 3: 14, 4: 12}
    font_size = font_sizes.get(level, 12)
    add_text_with_style(heading, text, bold=True, font_size=font_size)
    return heading

def process_text_line(paragraph, text):
    if '**' in text:
        parts = re.split(r'(\*\*[^*]+\*\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                add_text_with_style(paragraph, part[2:-2], bold=True)
            elif part.strip():
                add_text_with_style(paragraph, part)
    else:
        add_text_with_style(paragraph, text)

def main():
    print("🏆 最终版 Word 公式构建器")
    doc = create_base_template()
    
    chapters = [
        ("第一章-绪论.md", 1),
        ("第二章-相关背景与理论基础.md", 2),
        ("第三章-基于布隆过滤器的SQL和CTE动态缓存技术.md", 3),
        ("第四章-动态缓存更新技术与持久化技术.md", 4),
        ("第五章-实验与分析.md", 5),
        ("第六章-总结与展望.md", 6),
        ("统一参考文献列表.md", 7)
    ]
    
    for filename, chapter_num in chapters:
        filepath = BASE_DIR / filename
        if not filepath.exists():
            print(f"⚠️ 文件不存在: {filename}")
            continue
        
        print(f"\n📖 处理章节: {filename}")
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read()
        
        process_chapter(content, doc, chapter_num)
        
        if chapter_num < 7:
            doc.add_page_break()
            
    output_path = BASE_DIR / "final_thesis.docx"
    doc.save(str(output_path))
    print(f"\n✅ 文档已保存: {output_path}")

if __name__ == "__main__":
    main()