#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
改进的文档构建器
结合 enhanced_docx_builder.py 的公式处理和 native_equation_builder.py 的字体管理
解决合并文档时的字体和图片问题
"""

import os
import re
import json
import shutil
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.shared import OxmlElement, qn
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

# 导入辅助模块
from .improved_font_manager import force_font_style, add_text_with_style
from .improved_formula_processor import create_math_equation
from .improved_image_manager import copy_images_to_output, get_image_for_figure

BASE_DIR = Path(__file__).parent.absolute()
IMAGES_DIR = BASE_DIR / "images"
OUTPUT_IMAGES_DIR = BASE_DIR / "output_images"

def create_base_template():
    """创建基础模板文档"""
    doc = Document()
    
    # 删除默认段落
    for paragraph in doc.paragraphs:
        p = paragraph._element
        p.getparent().remove(p)
    
    # 设置文档默认字体
    doc.styles['Normal'].font.name = 'Times New Roman'
    doc.styles['Normal'].font.size = Pt(12)
    
    return doc

def add_heading_with_style(doc, text, level=1):
    """添加带样式的标题"""
    heading = doc.add_heading('', level=level)
    
    # 清空默认内容
    heading.clear()
    
    # 设置字体大小
    font_sizes = {1: 18, 2: 16, 3: 14, 4: 12}
    font_size = font_sizes.get(level, 12)
    
    # 添加文本
    add_text_with_style(heading, text, bold=True, font_size=font_size)
    
    return heading

def process_text_line(paragraph, text):
    """处理包含格式化的文本行"""
    if '**' in text:
        # 处理粗体
        parts = re.split(r'(\*\*[^*]+\*\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                add_text_with_style(paragraph, part[2:-2], bold=True)
            else:
                if part.strip():
                    add_text_with_style(paragraph, part)
    else:
        add_text_with_style(paragraph, text)

def main():
    """主函数"""
    print("🚀 改进的 Word 文档构建器")
    print("📝 特性：强制字体设置 + 改进公式处理 + 图片管理")
    
    # 复制图片到输出目录
    copy_images_to_output()
    
    # 创建基础模板
    doc = create_base_template()
    
    # 章节文件列表
    chapters = [
        ("第0章-摘要.md", 0),
        ("第一章-绪论.md", 1),
        ("第二章-相关背景与理论基础.md", 2),
        ("第三章-基于布隆过滤器的SQL和CTE动态缓存技术.md", 3),
        ("第四章-动态缓存更新技术与持久化技术.md", 4),
        ("第五章-实验与分析.md", 5),
        ("第六章-总结与展望.md", 6),
        ("统一参考文献列表.md", 7)
    ]
    
    # 导入章节处理器
    from .improved_chapter_processor import process_chapter
    
    for filename, chapter_num in chapters:
        filepath = BASE_DIR / filename
        
        if not filepath.exists():
            print(f"⚠️  文件不存在: {filename}")
            continue
        
        print(f"\n📖 处理章节: {filename}")
        
        # 读取文件内容
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read()
        
        # 处理内容
        process_chapter(content, doc, chapter_num)
        
        # 在章节之间添加分页符（除了最后一章）
        if chapter_num < 7:
            doc.add_page_break()
    
    # 保存文档
    output_path = BASE_DIR / "improved_thesis.docx"
    doc.save(str(output_path))
    
    print(f"\n✅ 文档已保存: {output_path}")
    
    # 验证文档
    try:
        test_doc = Document(str(output_path))
        
        # 统计图片数量
        image_count = 0
        for rel in test_doc.part.rels.values():
            if "image" in rel.target_ref:
                image_count += 1
        
        print(f"📊 文档统计:")
        print(f"   - 段落数量: {len(test_doc.paragraphs)}")
        print(f"   - 嵌入图片: {image_count}")
        print(f"   - 输出图片目录: {OUTPUT_IMAGES_DIR}")
        
        print("🏆 改进特性:")
        print("   - ✅ 强制字体设置 (Times New Roman)")
        print("   - ✅ 原生公式对象 + Unicode 备用")
        print("   - ✅ 图片复制到输出目录")
        print("   - ✅ 支持 PNG 和 Mermaid 文件")
        print("   - ✅ 统一的字体管理")
        
    except Exception as e:
        print(f"❌ 验证失败: {e}")

if __name__ == "__main__":
    main()