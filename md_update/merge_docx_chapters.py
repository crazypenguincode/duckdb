#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
合并 md_update 目录下的章节 DOCX 为一个完整的 DOCX，参考文献放在最后。
依赖: python-docx
安装: pip install python-docx

用法:
  cd /Users/max/src/duckdb/md_update
  python3 merge_docx_chapters.py

默认输入文件名（按顺序）:
  1. 第一章-绪论.docx
  2. 第二章-相关背景与理论基础.docx
  3. 第三章-动态缓存管理.docx
  4. 第四章-动态缓存更新技术与持久化技术.docx
  5. 第五章-实验与分析.docx
  6. 第六章-总结与展望.docx
  7. 统一参考文献列表.docx   (放在最后)

输出:
  合并输出为 merged_thesis.docx
"""

import os
import sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

CHAPTER_FILES = [
    "第一章-绪论.docx",
    "第二章-相关背景与理论基础.docx",
    "第三章-动态缓存管理.docx",
    "第四章-动态缓存更新技术与持久化技术.docx",
    "第五章-实验与分析.docx",
    "第六章-总结与展望.docx",
]

REFERENCE_FILE = "统一参考文献列表.docx"
OUTPUT_FILE = "merged_thesis.docx"


def add_page_break(doc: Document):
    p = doc.add_paragraph()
    run = p.add_run()
    # 直接插入分页符
    run.add_break(WD_BREAK.PAGE)


def append_document(dst: Document, src: Document, insert_page_break_before: bool = True):
    """
    将 src 的内容（段落、表格）追加到 dst 文档中，尽量保留格式。
    简化策略：逐段复制文本和基本格式，表格按单元格文本复制。
    """
    if insert_page_break_before and (len(dst.paragraphs) > 0 or len(dst.tables) > 0):
        add_page_break(dst)

    # 复制段落
    for para in src.paragraphs:
        new_para = dst.add_paragraph()
        # 复制段落样式名（若目标无对应样式，则回退默认）
        try:
            if para.style and para.style.name:
                new_para.style = para.style
        except Exception:
            pass
        for run in para.runs:
            new_run = new_para.add_run(run.text)
            # 基本格式迁移
            new_run.bold = run.bold
            new_run.italic = run.italic
            new_run.underline = run.underline
            new_run.font.name = getattr(run.font, "name", None)
            new_run.font.size = getattr(run.font, "size", None)
            new_run.font.color.rgb = getattr(getattr(run.font, "color", None), "rgb", None)

    # 复制表格（仅复制文本）
    for table in src.tables:
        new_table = dst.add_table(rows=len(table.rows), cols=len(table.columns))
        try:
            new_table.style = table.style
        except Exception:
            pass
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                new_table.cell(r_idx, c_idx).text = cell.text


def load_docx_safe(path: str):
    if not os.path.exists(path):
        print(f"[WARN] 文件缺失: {path}")
        return None
    try:
        return Document(path)
    except Exception as e:
        print(f"[ERROR] 打开文档失败: {path} -> {e}")
        return None


def main():
    os.chdir(BASE_DIR)
    print("工作目录:", BASE_DIR)

    out_path = os.path.join(BASE_DIR, OUTPUT_FILE)
    # 如果存在则覆盖
    if os.path.exists(out_path):
        try:
            os.remove(out_path)
        except Exception as e:
            print(f"[ERROR] 无法删除旧输出文件: {out_path} -> {e}")
            sys.exit(1)

    merged = Document()

    print("开始合并章节...")
    for fname in CHAPTER_FILES:
        path = os.path.join(BASE_DIR, fname)
        print(f"处理章节: {fname}")
        doc = load_docx_safe(path)
        if doc is None:
            continue
        # 章节标题（可选）
        # title_para = merged.add_paragraph(fname.replace(".docx", ""))
        # title_para.style = merged.styles['Heading 1'] if 'Heading 1' in [s.name for s in merged.styles] else None
        append_document(merged, doc, insert_page_break_before=(len(merged.paragraphs) > 0))

    # 参考文献
    ref_path = os.path.join(BASE_DIR, REFERENCE_FILE)
    print(f"处理参考文献: {REFERENCE_FILE}")
    ref_doc = load_docx_safe(ref_path)
    if ref_doc is not None:
        # 插入分页
        if len(merged.paragraphs) > 0:
            add_page_break(merged)
        # 插入“参考文献”标题
        heading = merged.add_paragraph("参考文献")
        # 设置为标题样式（若存在）
        try:
            heading.style = merged.styles['Heading 1']
        except Exception:
            pass
        # 再追加参考文献内容
        append_document(merged, ref_doc, insert_page_break_before=False)

    merged.save(out_path)
    print("合并完成。输出文件:", out_path)


if __name__ == "__main__":
    main()