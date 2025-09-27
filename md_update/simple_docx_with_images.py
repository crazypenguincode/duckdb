#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
简化版 DOCX 构建器，使用图片索引文件
"""

import os
import re
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

BASE_DIR = Path(__file__).parent.absolute()
IMAGES_DIR = BASE_DIR / "images"

def load_image_index():
    """加载图片索引"""
    index_file = BASE_DIR / "图片提取报告.json"
    if index_file.exists():
        with open(index_file, 'r', encoding='utf-8') as f:
            return json.load(f)
    return {}

def get_image_for_figure(fig_num, image_index):
    """根据图号获取对应的图片文件"""
    # 直接查找文件
    pattern = f"{fig_num}_*.png"
    matches = list(IMAGES_DIR.glob(pattern))
    if matches:
        return matches[0]
    
    # 如果没找到，尝试从索引中查找
    for chapter_data in image_index.get('chapters', {}).values():
        for img_info in chapter_data.get('images', []):
            if img_info.get('figure_number') == fig_num:
                img_path = IMAGES_DIR / img_info.get('png_file', '')
                if img_path.exists():
                    return img_path
    
    return None

def process_chapter(content, doc, chapter_num, image_index):
    """处理单个章节"""
    lines = content.split('\n')
    i = 0
    
    while i < len(lines):
        line = lines[i].strip()
        
        # 处理标题
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            title_text = line.lstrip('#').strip()
            
            if level <= 4:
                doc.add_heading(title_text, level=level)
            else:
                p = doc.add_paragraph()
                run = p.add_run(title_text)
                run.bold = True
        
        # 处理图片标题（寻找 **图X.Y 格式）
        elif line.startswith('**图') and line.endswith('**'):
            # 提取图号和标题
            match = re.search(r'\*\*图(\d+\.\d+)\s*(.*?)\*\*', line)
            if match:
                fig_num = match.group(1)
                fig_title = match.group(2).strip()
                
                print(f"处理图片: 图{fig_num} {fig_title}")
                
                # 查找对应的图片文件
                image_path = get_image_for_figure(fig_num, image_index)
                
                if image_path and image_path.exists():
                    try:
                        # 添加图片
                        doc.add_picture(str(image_path), width=Inches(5.5))
                        
                        # 添加图片标题
                        caption_p = doc.add_paragraph()
                        caption_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                        caption_run = caption_p.add_run(f"图{fig_num} {fig_title}")
                        caption_run.bold = True
                        
                        print(f"✅ 成功插入图片: {image_path.name}")
                        
                    except Exception as e:
                        print(f"❌ 插入图片失败: {e}")
                        # 添加占位符
                        p = doc.add_paragraph(f"[图片占位符: 图{fig_num} {fig_title}]")
                        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                else:
                    print(f"❌ 未找到图片文件: 图{fig_num}")
                    # 添加占位符
                    p = doc.add_paragraph(f"[图片缺失: 图{fig_num} {fig_title}]")
                    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        # 跳过 mermaid 代码块
        elif line.startswith('```mermaid'):
            j = i + 1
            while j < len(lines) and not lines[j].strip().startswith('```'):
                j += 1
            i = j  # 跳过整个代码块
        
        # 跳过其他代码块
        elif line.startswith('```'):
            j = i + 1
            while j < len(lines) and not lines[j].strip().startswith('```'):
                j += 1
            i = j
        
        # 处理普通段落
        elif line:
            # 处理内联公式 $...$
            if '$' in line and not line.startswith('$$'):
                p = doc.add_paragraph()
                # 简单处理：将公式标记为斜体
                formula_pattern = r'\$([^$]+)\$'
                parts = re.split(formula_pattern, line)
                
                for idx, part in enumerate(parts):
                    if idx % 2 == 1:  # 奇数索引是公式内容
                        run = p.add_run(part)
                        run.italic = True
                        run.font.name = 'Cambria Math'
                    else:  # 偶数索引是普通文本
                        if part.strip():
                            # 处理粗体
                            if '**' in part:
                                bold_parts = re.split(r'(\*\*[^*]+\*\*)', part)
                                for bold_part in bold_parts:
                                    if bold_part.startswith('**') and bold_part.endswith('**'):
                                        run = p.add_run(bold_part[2:-2])
                                        run.bold = True
                                    else:
                                        p.add_run(bold_part)
                            else:
                                p.add_run(part)
            
            # 处理块级公式 $$...$$
            elif line.startswith('$$') and line.endswith('$$'):
                formula = line[2:-2].strip()
                p = doc.add_paragraph()
                p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                run = p.add_run(formula)
                run.italic = True
                run.font.name = 'Cambria Math'
                run.font.size = Pt(12)
            
            # 普通段落
            else:
                if '**' in line:
                    p = doc.add_paragraph()
                    bold_parts = re.split(r'(\*\*[^*]+\*\*)', line)
                    for bold_part in bold_parts:
                        if bold_part.startswith('**') and bold_part.endswith('**'):
                            run = p.add_run(bold_part[2:-2])
                            run.bold = True
                        else:
                            p.add_run(bold_part)
                else:
                    doc.add_paragraph(line)
        
        i += 1

def main():
    """主函数"""
    print("🔧 简化版 DOCX 构建器")
    
    # 加载图片索引
    image_index = load_image_index()
    print(f"📊 加载图片索引: {len(image_index.get('chapters', {}))} 个章节")
    
    # 创建新文档
    doc = Document()
    
    # 章节文件列表
    chapters = [
        ("第一章-绪论.md", 1),
        ("第二章-相关背景与理论基础.md", 2),
        ("第三章-基于布隆过滤器的SQL和CTE动态缓存技术.md", 3),
        ("第四章-动态缓存更新技术与持久化技术.md", 4),
        ("第五章-实验与分析.md", 5),
        ("第六章-总结与展望.md", 6),
        ("统一参考文献列表.md", 7)
    ]
    
    total_images = 0
    
    for filename, chapter_num in chapters:
        filepath = BASE_DIR / filename
        
        if not filepath.exists():
            print(f"⚠️  文件不存在: {filename}")
            continue
        
        print(f"\n📖 处理章节: {filename}")
        
        # 读取文件内容
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read()
        
        # 处理内容
        process_chapter(content, doc, chapter_num, image_index)
        
        # 在章节之间添加分页符（除了最后一章）
        if chapter_num < 7:
            doc.add_page_break()
    
    # 保存文档
    output_path = BASE_DIR / "simple_built_thesis.docx"
    doc.save(str(output_path))
    
    print(f"\n✅ 文档已保存: {output_path}")
    
    # 验证文档
    try:
        test_doc = Document(str(output_path))
        
        # 统计图片数量
        image_count = 0
        for rel in test_doc.part.rels.values():
            if "image" in rel.target_ref:
                image_count += 1
        
        print(f"📊 文档统计:")
        print(f"   - 段落数量: {len(test_doc.paragraphs)}")
        print(f"   - 嵌入图片: {image_count}")
        
        if image_count > 0:
            print("🎉 成功插入图片！")
        else:
            print("⚠️  没有插入任何图片")
        
    except Exception as e:
        print(f"❌ 验证失败: {e}")

if __name__ == "__main__":
    main()