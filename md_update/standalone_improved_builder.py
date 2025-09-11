#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
独立版改进的文档构建器
解决合并文档时的字体和图片问题
结合两个原始文件的优点
"""

import os
import re
import json
import shutil
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.shared import OxmlElement, qn
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

BASE_DIR = Path(__file__).parent.absolute()
IMAGES_DIR = BASE_DIR / "images"
OUTPUT_IMAGES_DIR = BASE_DIR / "output_images"

def ensure_output_images_dir():
    """确保输出图片目录存在"""
    OUTPUT_IMAGES_DIR.mkdir(exist_ok=True)
    print(f"📁 输出图片目录: {OUTPUT_IMAGES_DIR}")

def copy_images_to_output():
    """复制图片到输出目录，支持 png 和 mmd 格式"""
    if not IMAGES_DIR.exists():
        print("⚠️  源图片目录不存在")
        return
    
    ensure_output_images_dir()
    
    # 复制所有 PNG 文件
    png_files = list(IMAGES_DIR.glob("*.png"))
    for png_file in png_files:
        dest_path = OUTPUT_IMAGES_DIR / png_file.name
        shutil.copy2(png_file, dest_path)
        print(f"📋 复制图片: {png_file.name}")
    
    # 查找并复制 mermaid 文件（如果存在）
    mmd_files = list(IMAGES_DIR.glob("*.mmd")) + list(IMAGES_DIR.glob("*.mermaid"))
    for mmd_file in mmd_files:
        dest_path = OUTPUT_IMAGES_DIR / mmd_file.name
        shutil.copy2(mmd_file, dest_path)
        print(f"📋 复制 Mermaid 文件: {mmd_file.name}")
    
    print(f"✅ 图片复制完成，共 {len(png_files)} 个 PNG 文件，{len(mmd_files)} 个 Mermaid 文件")

def force_font_style(run, font_name='Times New Roman', font_size=12, bold=False, italic=False):
    """强制设置字体样式 - 来自 native_equation_builder.py"""
    # 设置基本属性
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    
    # 获取或创建 rPr 元素
    rPr = run._element.get_or_add_rPr()
    
    # 清除现有字体设置
    for rFonts in rPr.xpath('.//w:rFonts'):
        rPr.remove(rFonts)
    
    # 强制设置字体
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    rPr.insert(0, rFonts)
    
    # 强制设置字体大小
    for sz in rPr.xpath('.//w:sz'):
        rPr.remove(sz)
    for szCs in rPr.xpath('.//w:szCs'):
        rPr.remove(szCs)
    
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(font_size * 2))  # Word 使用半点
    rPr.append(sz)
    
    szCs = OxmlElement('w:szCs')
    szCs.set(qn('w:val'), str(font_size * 2))
    rPr.append(szCs)

def add_text_with_style(paragraph, text, bold=False, italic=False, font_size=12):
    """添加带样式的文本"""
    run = paragraph.add_run(text)
    force_font_style(run, 'Times New Roman', font_size, bold, italic)
    return run

def create_math_equation(paragraph, latex_formula):
    """创建数学公式 - 改进版本，结合两种方法"""
    try:
        # 方法1：尝试使用简化的 Word 原生公式
        # 清理公式内容
        clean_formula = latex_formula.replace('\\', '').replace('{', '').replace('}', '')
        
        # 创建简单的数学元素
        math_xml = f"""
        <m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
            <m:r>
                <m:rPr>
                    <m:sty m:val="p"/>
                </m:rPr>
                <m:t>{clean_formula}</m:t>
            </m:r>
        </m:oMath>
        """
        
        # 尝试创建数学元素
        math_element = parse_xml(math_xml)
        paragraph._element.append(math_element)
        
        print(f"✅ 使用原生公式对象: {latex_formula[:30]}...")
        return True
        
    except Exception as e:
        print(f"⚠️  原生公式失败，使用 Unicode 替代: {str(e)[:50]}...")
        
        # 方法2：使用 Unicode 符号替代
        formula = process_unicode_formula(latex_formula)
        
        # 使用 Cambria Math 字体显示
        run = paragraph.add_run(formula)
        force_font_style(run, 'Cambria Math', 12, False, False)
        
        return False

def process_unicode_formula(latex_formula):
    """处理 Unicode 公式转换"""
    formula = latex_formula.strip()
    
    # 符号映射
    superscript_map = {
        '0': '⁰', '1': '¹', '2': '²', '3': '³', '4': '⁴',
        '5': '⁵', '6': '⁶', '7': '⁷', '8': '⁸', '9': '⁹',
        '+': '⁺', '-': '⁻', '=': '⁼', '(': '⁽', ')': '⁾',
        'n': 'ⁿ', 'i': 'ⁱ', 'x': 'ˣ', 'y': 'ʸ', 'k': 'ᵏ', 'm': 'ᵐ'
    }
    subscript_map = {
        '0': '₀', '1': '₁', '2': '₂', '3': '₃', '4': '₄',
        '5': '₅', '6': '₆', '7': '₇', '8': '₈', '9': '₉',
        '+': '₊', '-': '₋', '=': '₌', '(': '₍', ')': '₎',
        'a': 'ₐ', 'e': 'ₑ', 'h': 'ₕ', 'i': 'ᵢ', 'j': 'ⱼ',
        'k': 'ₖ', 'l': 'ₗ', 'm': 'ₘ', 'n': 'ₙ', 'o': 'ₒ',
        'p': 'ₚ', 'r': 'ᵣ', 's': 'ₛ', 't': 'ₜ', 'u': 'ᵤ',
        'v': 'ᵥ', 'x': 'ₓ'
    }
    symbol_map = {
        r'\\alpha': 'α', r'\\beta': 'β', r'\\gamma': 'γ', r'\\delta': 'δ',
        r'\\epsilon': 'ε', r'\\theta': 'θ', r'\\lambda': 'λ', r'\\mu': 'μ',
        r'\\pi': 'π', r'\\sigma': 'σ', r'\\phi': 'φ', r'\\omega': 'ω',
        r'\\infty': '∞', r'\\to': '→', r'\\leq': '≤', r'\\geq': '≥',
        r'\\neq': '≠', r'\\approx': '≈', r'\\cdot': '·', r'\\times': '×',
        r'\\div': '÷', r'\\pm': '±'
    }
    
    def can_be_unicode(text, conversion_map):
        return all(char in conversion_map for char in text)

    def to_unicode(text, conversion_map):
        return "".join(conversion_map.get(char, char) for char in text)

    def replace_superscript(match):
        base, sup = match.groups()
        if sup.startswith('{') and sup.endswith('}'):
            sup = sup[1:-1]
        
        if can_be_unicode(sup, superscript_map):
            return f'{base}{to_unicode(sup, superscript_map)}'
        else:
            return f'{base}^{sup}'

    def replace_subscript(match):
        base, sub = match.groups()
        if sub.startswith('{') and sub.endswith('}'):
            sub = sub[1:-1]
        
        if sub in ['opt', 'min', 'max'] or not can_be_unicode(sub, subscript_map):
            return f'{base}_{sub}'
        else:
            return f'{base}{to_unicode(sub, subscript_map)}'
    
    # 处理公式
    previous_formula = ""
    while formula != previous_formula:
        previous_formula = formula
        
        # 移除 \left 和 \right
        formula = formula.replace('\\left', '').replace('\\right', '')
        
        # 处理函数
        formula = re.sub(r'\\ln\b', 'ln', formula)
        formula = re.sub(r'\\log\b', 'log', formula)
        
        # 处理分数
        formula = re.sub(r'\\frac\{d\}\{d([a-zA-Z])\}', r'd/d\1', formula)
        formula = re.sub(r'\\frac\{([^}]+)\}\{([^}]+)\}', r'(\1)/(\2)', formula)
        
        # 处理上下标
        formula = re.sub(r'(\w+|\([^)]+\))\^(\{[^}]+\}|[^{}\s^]+)', replace_superscript, formula)
        formula = re.sub(r'(\w+|\([^)]+\))_(\{[^}]+\}|[^{}\s_]+)', replace_subscript, formula)
        
        # 处理根号
        formula = re.sub(r'\\sqrt\{([^}]+)\}', r'√(\1)', formula)
        
        # 处理符号
        for latex_symbol, unicode_symbol in symbol_map.items():
            formula = re.sub(latex_symbol, unicode_symbol, formula)
    
    return formula

def get_image_for_figure(fig_num):
    """根据图号获取对应的图片文件"""
    # 首先在输出目录查找
    pattern = f"{fig_num}_*.png"
    matches = list(OUTPUT_IMAGES_DIR.glob(pattern))
    if matches:
        return matches[0]
    
    # 如果输出目录没有，在原始目录查找
    matches = list(IMAGES_DIR.glob(pattern))
    if matches:
        return matches[0]
    
    return None

def create_base_template():
    """创建基础模板文档"""
    doc = Document()
    
    # 删除默认段落
    for paragraph in doc.paragraphs:
        p = paragraph._element
        p.getparent().remove(p)
    
    # 设置文档默认字体
    doc.styles['Normal'].font.name = 'Times New Roman'
    doc.styles['Normal'].font.size = Pt(12)
    
    return doc

def add_heading_with_style(doc, text, level=1):
    """添加带样式的标题"""
    heading = doc.add_heading('', level=level)
    
    # 清空默认内容
    heading.clear()
    
    # 设置字体大小
    font_sizes = {1: 18, 2: 16, 3: 14, 4: 12}
    font_size = font_sizes.get(level, 12)
    
    # 添加文本
    add_text_with_style(heading, text, bold=True, font_size=font_size)
    
    return heading

def process_text_line(paragraph, text):
    """处理包含格式化的文本行"""
    if '**' in text:
        # 处理粗体
        parts = re.split(r'(\*\*[^*]+\*\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                add_text_with_style(paragraph, part[2:-2], bold=True)
            else:
                if part.strip():
                    add_text_with_style(paragraph, part)
    else:
        add_text_with_style(paragraph, text)

def process_chapter(content, doc, chapter_num):
    """处理单个章节"""
    lines = content.split('\n')
    i = 0
    formula_count = 0
    native_formula_count = 0
    
    while i < len(lines):
        line = lines[i].strip()
        
        # 处理标题
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            title_text = line.lstrip('#').strip()
            
            if level <= 4:
                add_heading_with_style(doc, title_text, level)
            else:
                p = doc.add_paragraph()
                add_text_with_style(p, title_text, bold=True)
        
        # 处理图片标题
        elif line.startswith('**图') and line.endswith('**'):
            match = re.search(r'\*\*图(\d+\.\d+)\s*(.*?)\*\*', line)
            if match:
                fig_num = match.group(1)
                fig_title = match.group(2).strip()
                
                print(f"处理图片: 图{fig_num} {fig_title}")
                
                image_path = get_image_for_figure(fig_num)
                
                if image_path and image_path.exists():
                    try:
                        # 添加图片
                        doc.add_picture(str(image_path), width=Inches(5.5))
                        
                        # 添加图片标题
                        caption_p = doc.add_paragraph()
                        caption_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                        add_text_with_style(caption_p, f"图{fig_num} {fig_title}", bold=True)
                        
                        print(f"✅ 成功插入图片: {image_path.name}")
                        
                    except Exception as e:
                        print(f"❌ 插入图片失败: {e}")
                        p = doc.add_paragraph()
                        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                        add_text_with_style(p, f"[图片占位符: 图{fig_num} {fig_title}]")
                else:
                    print(f"❌ 未找到图片文件: 图{fig_num}")
                    p = doc.add_paragraph()
                    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                    add_text_with_style(p, f"[图片缺失: 图{fig_num} {fig_title}]")
        
        # 跳过 mermaid 代码块
        elif line.startswith('```mermaid'):
            j = i + 1
            while j < len(lines) and not lines[j].strip().startswith('```'):
                j += 1
            i = j
        
        # 处理其他代码块
        elif line.startswith('```'):
            j = i + 1
            code_lines = []
            while j < len(lines) and not lines[j].strip().startswith('```'):
                code_lines.append(lines[j])
                j += 1
            
            # 添加代码块
            if code_lines:
                code_text = '\n'.join(code_lines)
                p = doc.add_paragraph()
                run = p.add_run(code_text)
                force_font_style(run, 'Consolas', 10)
            
            i = j
        
        # 处理普通段落
        elif line:
            # 处理内联公式 $...$
            if '$' in line and not line.startswith('$$'):
                p = doc.add_paragraph()
                formula_pattern = r'\$([^$]+)\$'
                parts = re.split(formula_pattern, line)
                
                for idx, part in enumerate(parts):
                    if idx % 2 == 1:  # 奇数索引是公式内容
                        if create_math_equation(p, part):
                            native_formula_count += 1
                        formula_count += 1
                    else:  # 偶数索引是普通文本
                        if part.strip():
                            process_text_line(p, part)
            
            # 处理块级公式 $$...$$
            elif line.startswith('$$') and line.endswith('$$'):
                formula = line[2:-2].strip()
                p = doc.add_paragraph()
                p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                if create_math_equation(p, formula):
                    native_formula_count += 1
                formula_count += 1
            
            # 处理多行块级公式
            elif line.startswith('$$'):
                formula_lines = [line[2:]]  # 去掉开头的 $$
                j = i + 1
                while j < len(lines):
                    if lines[j].strip().endswith('$$'):
                        formula_lines.append(lines[j].strip()[:-2])  # 去掉结尾的 $$
                        break
                    else:
                        formula_lines.append(lines[j])
                    j += 1
                
                formula = '\n'.join(formula_lines).strip()
                p = doc.add_paragraph()
                p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                if create_math_equation(p, formula):
                    native_formula_count += 1
                formula_count += 1
                i = j
            
            # 普通段落
            else:
                p = doc.add_paragraph()
                process_text_line(p, line)
        
        i += 1
    
    print(f"📊 章节 {chapter_num} 公式统计: {formula_count} 个公式，{native_formula_count} 个原生公式")

def main():
    """主函数"""
    print("🚀 独立版改进的 Word 文档构建器")
    print("📝 特性：强制字体设置 + 改进公式处理 + 图片管理")
    
    # 复制图片到输出目录
    copy_images_to_output()
    
    # 创建基础模板
    doc = create_base_template()
    
    # 章节文件列表
    chapters = [
        ("第0章-摘要.md", 0),
        ("第一章-绪论.md", 1),
        ("第二章-相关背景与理论基础.md", 2),
        ("第三章-动态缓存管理.md", 3),
        ("第四章-动态缓存更新技术与持久化技术.md", 4),
        ("第五章-实验与分析.md", 5),
        ("第六章-总结与展望.md", 6),
        ("统一参考文献列表.md", 7)
    ]
    
    for filename, chapter_num in chapters:
        filepath = BASE_DIR / filename
        
        if not filepath.exists():
            print(f"⚠️  文件不存在: {filename}")
            continue
        
        print(f"\n📖 处理章节: {filename}")
        
        # 读取文件内容
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read()
        
        # 处理内容
        process_chapter(content, doc, chapter_num)
        
        # 在章节之间添加分页符（除了最后一章）
        if chapter_num < 7:
            doc.add_page_break()
    
    # 保存文档
    output_path = BASE_DIR / "standalone_improved_thesis.docx"
    doc.save(str(output_path))
    
    print(f"\n✅ 文档已保存: {output_path}")
    
    # 验证文档
    try:
        test_doc = Document(str(output_path))
        
        # 统计图片数量
        image_count = 0
        for rel in test_doc.part.rels.values():
            if "image" in rel.target_ref:
                image_count += 1
        
        print(f"📊 文档统计:")
        print(f"   - 段落数量: {len(test_doc.paragraphs)}")
        print(f"   - 嵌入图片: {image_count}")
        print(f"   - 输出图片目录: {OUTPUT_IMAGES_DIR}")
        
        print("🏆 改进特性:")
        print("   - ✅ 强制字体设置 (Times New Roman)")
        print("   - ✅ 简化原生公式 + Unicode 备用")
        print("   - ✅ 图片复制到输出目录")
        print("   - ✅ 支持 PNG 和 Mermaid 文件")
        print("   - ✅ 统一的字体管理")
        print("   - ✅ 解决合并文档问题")
        
    except Exception as e:
        print(f"❌ 验证失败: {e}")

if __name__ == "__main__":
    main()