#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
终极版 DOCX 构建器：确保图片和公式正确显示
1) 使用绝对路径引用图片
2) 验证所有文件路径
3) 测试多种 pandoc 参数组合
4) 详细的调试和验证
"""

import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

BASE_DIR = Path(__file__).parent.absolute()
IMAGES_DIR = BASE_DIR / "images"
BUILD_DIR = BASE_DIR / "_ultimate_build"
BUILD_MD_DIR = BUILD_DIR / "md"
DOCX_OUT_DIR = BUILD_DIR / "docx"
MERGED_DOCX = BASE_DIR / "ultimate_merged_thesis.docx"

CHAPTERS = [
    "第一章-绪论.md",
    "第二章-相关背景与理论基础.md", 
    "第三章-动态缓存管理.md",
    "第四章-动态缓存更新技术与持久化技术.md",
    "第五章-实验与分析.md",
    "第六章-总结与展望.md",
]
REF_MD = "统一参考文献列表.md"

def run_cmd(cmd, cwd=None):
    print(f"+ {' '.join(str(x) for x in cmd)}")
    cp = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, 
                       text=True, cwd=cwd or BASE_DIR)
    print(f"  返回码: {cp.returncode}")
    if cp.stdout.strip():
        print(f"  输出: {cp.stdout.strip()}")
    if cp.stderr.strip():
        print(f"  错误: {cp.stderr.strip()}")
    return cp

def setup_build_environment():
    """设置构建环境"""
    print("=== 设置构建环境 ===")
    
    # 清理并创建构建目录
    if BUILD_DIR.exists():
        shutil.rmtree(BUILD_DIR)
    
    BUILD_DIR.mkdir(exist_ok=True)
    BUILD_MD_DIR.mkdir(exist_ok=True)
    DOCX_OUT_DIR.mkdir(exist_ok=True)
    
    print(f"构建目录: {BUILD_DIR}")
    print(f"图片目录: {IMAGES_DIR}")
    print(f"图片目录存在: {IMAGES_DIR.exists()}")
    
    if IMAGES_DIR.exists():
        png_files = list(IMAGES_DIR.glob("*.png"))
        print(f"找到 {len(png_files)} 个 PNG 文件")
        
        # 只显示前几个文件名作为示例
        for i, img in enumerate(png_files[:5]):
            print(f"  示例图片 {i+1}: {img.name}")
        if len(png_files) > 5:
            print(f"  ... 还有 {len(png_files) - 5} 个图片")

def get_chapter_images(chapter_num):
    """获取指定章节的图片列表"""
    if not IMAGES_DIR.exists():
        return []
    
    images = []
    for img in IMAGES_DIR.glob("*.png"):
        if img.name.startswith(f"{chapter_num}.") and not img.name.startswith("mermaid_"):
            images.append(img)
    
    return sorted(images, key=lambda x: x.name)

def replace_mermaid_with_images(md_text, chapter_num):
    """替换 mermaid 代码块为图片引用"""
    chapter_images = get_chapter_images(chapter_num)
    
    # 查找所有 mermaid 代码块
    mermaid_pattern = r"```mermaid[^\n]*\n([\s\S]*?)```"
    blocks = list(re.finditer(mermaid_pattern, md_text))
    
    if not blocks:
        return md_text, 0
    
    print(f"  找到 {len(blocks)} 个 mermaid 块")
    print(f"  可用图片 {len(chapter_images)} 个:")
    for img in chapter_images:
        print(f"    - {img.name} (存在: {img.exists()})")
    
    # 从后往前替换，避免位置偏移
    replaced = 0
    for i, block in enumerate(reversed(blocks)):
        img_idx = len(blocks) - 1 - i
        if img_idx < len(chapter_images):
            img_path = chapter_images[img_idx]
            
            # 验证图片文件存在
            if not img_path.exists():
                print(f"    ❌ 图片文件不存在: {img_path}")
                continue
            
            # 使用绝对路径
            abs_img_path = str(img_path.absolute())
            replacement = f"\n![{img_path.stem}]({abs_img_path})\n\n"
            md_text = md_text[:block.start()] + replacement + md_text[block.end():]
            replaced += 1
            print(f"    ✅ 替换 mermaid 块 -> {img_path.name}")
            print(f"       绝对路径: {abs_img_path}")
        else:
            print(f"    ⚠️  跳过 mermaid 块（无对应图片）")
    
    return md_text, replaced

def fix_math_formulas(md_text):
    """修复数学公式格式"""
    print("  修复数学公式...")
    
    # 统计原始公式数量
    original_inline = len(re.findall(r"\$[^$\n]+\$", md_text))
    original_block = len(re.findall(r"\$\$[^$]+\$\$", md_text))
    print(f"    原始内联公式: {original_inline} 个")
    print(f"    原始块级公式: {original_block} 个")
    
    # 移除被反引号包裹的公式
    md_text = re.sub(r"`(\$[^`$]+\$)`", r"\1", md_text)
    
    # 确保行间公式独占一行，前后有空行
    md_text = re.sub(r"([^\n])\$\$([^$]+)\$\$([^\n])", r"\1\n\n$$\2$$\n\n\3", md_text)
    
    # 修复常见的公式问题
    md_text = re.sub(r"\\frac\s*\{\s*([^}]+)\s*\}\s*\{\s*([^}]+)\s*\}", r"\\frac{\1}{\2}", md_text)
    
    # 确保内联公式前后有空格（但不破坏句子结构）
    md_text = re.sub(r"([a-zA-Z0-9])\$([^$]+)\$([a-zA-Z0-9])", r"\1 $\2$ \3", md_text)
    
    # 统计处理后公式数量
    final_inline = len(re.findall(r"\$[^$\n]+\$", md_text))
    final_block = len(re.findall(r"\$\$[^$]+\$\$", md_text))
    print(f"    处理后内联公式: {final_inline} 个")
    print(f"    处理后块级公式: {final_block} 个")
    
    return md_text

def preprocess_markdown(md_path):
    """预处理单个 Markdown 文件"""
    print(f"\n=== 预处理: {md_path.name} ===")
    
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    print(f"原始文件大小: {len(content)} 字符")
    
    # 提取章节号
    chapter_match = re.search(r"第([一二三四五六七八九十\d]+)章", md_path.name)
    if chapter_match:
        chapter_text = chapter_match.group(1)
        chinese_to_num = {'一': '1', '二': '2', '三': '3', '四': '4', '五': '5', '六': '6'}
        chapter_num = chinese_to_num.get(chapter_text, chapter_text)
    else:
        chapter_num = "0"
    
    print(f"章节号: {chapter_num}")
    
    # 替换 mermaid
    content, mermaid_count = replace_mermaid_with_images(content, chapter_num)
    
    # 修复公式
    content = fix_math_formulas(content)
    
    print(f"处理后文件大小: {len(content)} 字符")
    print(f"替换了 {mermaid_count} 个 mermaid 图")
    
    return content, mermaid_count

def test_pandoc_conversion(md_path, docx_path):
    """测试不同的 pandoc 转换方式"""
    print(f"\n=== 转换测试: {md_path.name} ===")
    
    # 多种转换方式
    conversion_methods = [
        {
            "name": "标准转换 + 数学支持",
            "cmd": [
                "pandoc", str(md_path), "-o", str(docx_path),
                "--standalone",
                "--from", "markdown+tex_math_dollars",
                "--to", "docx"
            ]
        },
        {
            "name": "增强数学支持",
            "cmd": [
                "pandoc", str(md_path), "-o", str(docx_path),
                "--standalone",
                "--from", "markdown+tex_math_dollars+tex_math_double_backslash",
                "--to", "docx",
                "--mathml"
            ]
        },
        {
            "name": "保留格式",
            "cmd": [
                "pandoc", str(md_path), "-o", str(docx_path),
                "--standalone",
                "--from", "markdown+tex_math_dollars",
                "--to", "docx",
                "--wrap", "preserve"
            ]
        },
        {
            "name": "基础转换",
            "cmd": [
                "pandoc", str(md_path), "-o", str(docx_path),
                "--from", "markdown",
                "--to", "docx"
            ]
        }
    ]
    
    for i, method in enumerate(conversion_methods):
        print(f"\n--- 尝试方法 {i+1}: {method['name']} ---")
        
        # 删除之前的输出文件
        if docx_path.exists():
            docx_path.unlink()
        
        result = run_cmd(method['cmd'])
        
        if result.returncode == 0 and docx_path.exists():
            file_size = docx_path.stat().st_size
            print(f"✅ 转换成功！文件大小: {file_size} 字节")
            return True
        else:
            print(f"❌ 转换失败")
    
    print(f"❌ 所有转换方法都失败了")
    return False

def verify_docx_content(docx_path):
    """验证 DOCX 文件内容"""
    try:
        from docx import Document
        doc = Document(str(docx_path))
        
        print(f"\n=== 验证 DOCX: {docx_path.name} ===")
        print(f"段落数量: {len(doc.paragraphs)}")
        
        # 检查是否包含图片
        image_count = 0
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                image_count += 1
        
        print(f"嵌入图片数量: {image_count}")
        
        # 检查数学公式（查找 OMML 或 MathML）
        math_count = 0
        for paragraph in doc.paragraphs:
            if "math" in paragraph._element.xml.lower():
                math_count += 1
        
        print(f"数学公式段落: {math_count}")
        
        return True
        
    except Exception as e:
        print(f"❌ 验证失败: {e}")
        return False

def merge_docx_files(docx_files, output_path):
    """合并多个 DOCX 文件"""
    print(f"\n=== 合并 DOCX 文件 ===")
    
    try:
        from docx import Document
        from docx.enum.text import WD_BREAK
    except ImportError:
        print("[ERROR] 需要安装 python-docx: pip install python-docx")
        return False
    
    merged = Document()
    
    for i, docx_path in enumerate(docx_files):
        if not docx_path.exists():
            print(f"⚠️  跳过缺失文件: {docx_path}")
            continue
        
        print(f"合并文件 {i+1}: {docx_path.name}")
        
        try:
            doc = Document(str(docx_path))
        except Exception as e:
            print(f"❌ 无法打开文件 {docx_path}: {e}")
            continue
        
        # 添加分页符（除了第一个文档）
        if i > 0:
            merged.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        
        # 复制内容
        for element in doc.element.body:
            merged.element.body.append(element)
    
    merged.save(str(output_path))
    print(f"✅ 合并完成: {output_path}")
    
    # 验证合并后的文件
    verify_docx_content(output_path)
    
    return True

def main():
    print("🚀 终极版 DOCX 构建器启动")
    print(f"工作目录: {BASE_DIR}")
    
    # 设置构建环境
    setup_build_environment()
    
    docx_files = []
    
    # 处理所有章节
    all_files = CHAPTERS + [REF_MD]
    
    for md_name in all_files:
        md_path = BASE_DIR / md_name
        if not md_path.exists():
            print(f"⚠️  文件不存在: {md_name}")
            continue
        
        # 预处理
        processed_content, mermaid_count = preprocess_markdown(md_path)
        
        # 保存预处理后的文件
        processed_md = BUILD_MD_DIR / md_name
        with open(processed_md, 'w', encoding='utf-8') as f:
            f.write(processed_content)
        
        print(f"✅ 预处理完成，保存到: {processed_md}")
        
        # 转换为 DOCX
        docx_name = md_name.replace('.md', '.docx')
        docx_path = DOCX_OUT_DIR / docx_name
        
        if test_pandoc_conversion(processed_md, docx_path):
            # 验证转换结果
            if verify_docx_content(docx_path):
                docx_files.append(docx_path)
            else:
                print(f"❌ DOCX 验证失败: {docx_name}")
        else:
            print(f"❌ 转换失败: {md_name}")
    
    # 合并所有 DOCX
    if docx_files:
        print(f"\n📚 准备合并 {len(docx_files)} 个 DOCX 文件")
        if merge_docx_files(docx_files, MERGED_DOCX):
            print(f"\n🎉 大功告成！")
            print(f"📄 最终文档: {MERGED_DOCX}")
            print(f"📁 构建目录: {BUILD_DIR}")
            print(f"📝 预处理的 MD: {BUILD_MD_DIR}")
            print(f"📄 单独的 DOCX: {DOCX_OUT_DIR}")
        else:
            print(f"\n❌ 合并失败")
    else:
        print(f"\n❌ 没有成功转换的 DOCX 文件")

if __name__ == "__main__":
    main()